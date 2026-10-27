"""DOCX text extraction."""

from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from ._common import build_document
from .exceptions import InvalidDocument
from .models import ParsedDocument

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


class DocxParser:
    def parse(
        self,
        source: Path,
        *,
        filename: str,
        media_type: str,
        size_bytes: int,
    ) -> ParsedDocument:
        try:
            document = Document(source)
        except (BadZipFile, PackageNotFoundError, ValueError, OSError) as exc:
            raise InvalidDocument(f"Invalid DOCX document: {filename}") from exc

        blocks = [(paragraph.text, None) for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append((" | ".join(cell.text for cell in row.cells), None))

        return build_document(
            filename=filename,
            media_type=media_type or DOCX_MEDIA_TYPE,
            format="docx",
            size_bytes=size_bytes,
            blocks=blocks,
        )
